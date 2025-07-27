# metadatos.py

#-----
import os
import datetime
import pandas as pd
import openpyxl
import olefile
import docx
from PyPDF2 import PdfReader
from pptx import Presentation
from tkinter import filedialog, messagebox, ttk
import tkinter as tk
from hashlib import sha256

def calcular_hash(path):
    """Calcula y devuelve el SHA-256 de un archivo."""
    try:
        h = sha256()
        with open(path, 'rb') as f:
            while chunk := f.read(8192):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return "Error"

def extraer_metadatos(archivo):
    """Extrae metadatos según la extensión del archivo."""
    ext = os.path.splitext(archivo)[1].lower()
    try:
        if ext == '.xlsx':
            wb = openpyxl.load_workbook(archivo, data_only=True)
            p = wb.properties
            return [
                p.creator or "-", p.lastModifiedBy or "-",
                p.created.strftime("%Y-%m-%d %H:%M:%S") if p.created else "-",
                p.modified.strftime("%Y-%m-%d %H:%M:%S") if p.modified else "-"
            ]
        elif ext == '.xls':
            with olefile.OleFileIO(archivo) as ole:
                m = ole.get_metadata()
                return [
                    m.author or "-", m.last_saved_by or "-",
                    m.create_time.strftime("%Y-%m-%d %H:%M:%S") if m.create_time else "-",
                    m.last_saved_time.strftime("%Y-%m-%d %H:%M:%S") if m.last_saved_time else "-"
                ]
        elif ext == '.docx':
            d = docx.Document(archivo)
            p = d.core_properties
            return [
                p.author or "-", p.last_modified_by or "-",
                p.created.strftime("%Y-%m-%d %H:%M:%S") if p.created else "-",
                p.modified.strftime("%Y-%m-%d %H:%M:%S") if p.modified else "-"
            ]
        elif ext == '.doc':
            return ["No soportado"]*4
        elif ext == '.pdf':
            r = PdfReader(archivo)
            i = r.metadata
            # Metadata PDF puede venir en formato distinto, normalizo
            def fmt_pdf_date(d):
                if not d:
                    return "-"
                try:
                    # PDF date string ejemplo: D:20230405123000Z
                    return datetime.datetime.strptime(d[2:16], "%Y%m%d%H%M%S").strftime("%Y-%m-%d %H:%M:%S")
                except Exception:
                    return d
            return [
                i.author or "-",
                i.producer or "-",
                fmt_pdf_date(i.creation_date),
                fmt_pdf_date(i.modification_date)
            ]
        elif ext == '.pptx':
            pr = Presentation(archivo)
            p = pr.core_properties
            return [
                p.author or "-", p.last_modified_by or "-",
                p.created.strftime("%Y-%m-%d %H:%M:%S") if p.created else "-",
                p.modified.strftime("%Y-%m-%d %H:%M:%S") if p.modified else "-"
            ]
        elif ext == '.ppt':
            return ["No soportado"]*4
        else:
            # Para texto plano, csv y otros: fechas del sistema
            stat = os.stat(archivo)
            c = datetime.datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S")
            m = datetime.datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            return ["N/A","N/A",c,m]
    except Exception:
        return ["Error"]*4

def mostrar_metadatos(root):
    """
    Muestra en un Toplevel de Tkinter una tabla con
    [Archivo, Creador, Editor, Creación, Modificación, SHA-256].
    Ofrece guardar como CSV.
    """
    files = filedialog.askopenfilenames(
        filetypes=[("Soportados","*.xlsx *.xls *.xlsb *.csv *.txt *.py *.docx *.doc *.pdf *.pptx *.ppt")]
    )
    if not files:
        return

    win = tk.Toplevel(root)
    win.title("Metadatos y Hash")
    cols = ["Archivo","Creador","Editor","Creación","Modificación","SHA-256"]
    tv = ttk.Treeview(win, columns=cols, show='headings')
    for c in cols:
        tv.heading(c, text=c)
        tv.column(c, width=130, anchor="center")
    tv.pack(fill=tk.BOTH, expand=True)

    data = []
    for f in files:
        m = extraer_metadatos(f)
        h = calcular_hash(f)
        row = [os.path.basename(f), *m, h]
        tv.insert("", tk.END, values=row)
        data.append(row)

    if messagebox.askyesno("Guardar CSV","¿Guardar metadatos en CSV?"):
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV","*.csv")])
        if path:
            pd.DataFrame(data, columns=cols).to_csv(path, index=False, encoding="utf-8")
            messagebox.showinfo("Guardado", f"CSV guardado en {path}")
